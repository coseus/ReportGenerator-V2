# -*- coding: utf-8 -*-
"""Standalone Executive Conclusion (Concluzie executivă) — DOCX + PDF.

A short management-facing deliverable, exported separately from the full report.
Content comes from util.exec_conclusion (auto-filled from the report, with the
narrative parts editable from the web UI)."""

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)

from util import exec_conclusion as EC
from util.i18n import get_language
from report.pdf_generator import F_REG, F_BOLD  # registered DejaVu fonts


def _safe(v):
    return "" if v is None else str(v)


# =========================================================================
# DOCX
# =========================================================================
def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _center(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def generate_conclusion_docx_bytes(report: dict) -> bytes:
    lang = get_language(report)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)

    accent = (report.get("theme_hex") or "#2E3B4E").lstrip("#")

    # Title block
    t = _center(doc.add_paragraph()); r = t.add_run(EC.title(lang))
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor.from_string(accent)
    st = _center(doc.add_paragraph()); r = st.add_run(EC.subtitle(report, lang))
    r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    who = " · ".join([x for x in [_safe(report.get("client")).strip(), _safe(report.get("date")).strip()] if x])
    if who:
        p = _center(doc.add_paragraph()); rr = p.add_run(who); rr.font.size = Pt(10); rr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    doc.add_paragraph()
    for para in EC.intro(report, lang):
        doc.add_paragraph(para)

    # Metadata table
    rows = EC.metadata_rows(report, lang)
    tbl = doc.add_table(rows=0, cols=2); tbl.style = "Table Grid"
    for label, value in rows:
        cells = tbl.add_row().cells
        _shade(cells[0], "EEF2F7")
        rn = cells[0].paragraphs[0].add_run(label); rn.bold = True
        cells[1].text = _safe(value)
    _set_col_widths(tbl, [Inches(1.7), Inches(4.9)])

    doc.add_paragraph()
    # Risk banner
    risk_label, risk_color, risk_desc = EC.risk(report, lang)
    rt = doc.add_table(rows=0, cols=2); rt.style = "Table Grid"
    hdr = rt.add_row().cells
    _merge_row(hdr); _shade(hdr[0], accent)
    hr = hdr[0].paragraphs[0].add_run(EC.section("risk_general", lang)); hr.bold = True; hr.font.color.rgb = RGBColor(255, 255, 255)
    body = rt.add_row().cells
    _shade(body[0], risk_color)
    br = body[0].paragraphs[0].add_run(risk_label); br.bold = True; br.font.color.rgb = RGBColor(255, 255, 255); br.font.size = Pt(13)
    body[1].text = risk_desc
    _set_col_widths(rt, [Inches(1.7), Inches(4.9)])

    # Posture line
    pl = EC.posture_line(report, lang)
    if pl:
        pt = doc.add_table(rows=1, cols=1); pt.style = "Table Grid"
        _shade(pt.rows[0].cells[0], "F1F3F5")
        pr = pt.rows[0].cells[0].paragraphs[0].add_run(pl); pr.bold = True

    # Conclusion narrative
    _h1(doc, EC.section("conclusion", lang), accent)
    for para in EC.narrative(report, lang):
        doc.add_paragraph(para)

    # Impact table
    _h1(doc, EC.section("impact", lang), accent)
    ih = EC.impact_header(lang)
    it = doc.add_table(rows=1, cols=2); it.style = "Table Grid"
    for i, htxt in enumerate(ih):
        _shade(it.rows[0].cells[i], "EEF2F7")
        rr = it.rows[0].cells[i].paragraphs[0].add_run(htxt); rr.bold = True
    for area, impact in EC.impact_rows(lang):
        c = it.add_row().cells
        rr = c[0].paragraphs[0].add_run(area); rr.bold = True
        c[1].text = impact
    _set_col_widths(it, [Inches(1.9), Inches(4.7)])

    # Remediation directions
    _h1(doc, EC.section("remediation", lang), accent)
    doc.add_paragraph(EC.remediation_intro(report, lang))
    dh = EC.directions_header(lang)
    dt = doc.add_table(rows=1, cols=4); dt.style = "Table Grid"
    for i, htxt in enumerate(dh):
        _shade(dt.rows[0].cells[i], "EEF2F7")
        rr = dt.rows[0].cells[i].paragraphs[0].add_run(htxt); rr.bold = True
    for idx, d in enumerate(EC.directions(report, lang), start=1):
        c = dt.add_row().cells
        c[0].text = str(idx)
        c[1].text = _safe(d.get("direction"))
        c[2].text = _safe(d.get("intervention"))
        c[3].text = _safe(d.get("urgency"))
    _set_col_widths(dt, [Inches(0.4), Inches(3.0), Inches(2.4), Inches(0.9)])

    # Compliance — priority findings mapped to controls (table only)
    mrows = EC.compliance_mapping_rows(report, lang)
    if mrows:
        _h1(doc, EC.section("mapping", lang), accent)
        doc.add_paragraph(EC.mapping_lead(lang))
        mh = EC.mapping_header(lang)
        mt = doc.add_table(rows=1, cols=3); mt.style = "Table Grid"
        for i, h in enumerate(mh):
            _shade(mt.rows[0].cells[i], "EEF2F7")
            rr = mt.rows[0].cells[i].paragraphs[0].add_run(h); rr.bold = True
        for row in mrows:
            c = mt.add_row().cells
            c[0].text = _safe(row["title"])
            c[1].text = _safe(row["severity"])
            c[2].text = _safe(row["domain"])
        _set_col_widths(mt, [Inches(3.1), Inches(1.1), Inches(2.4)])

    bio = BytesIO(); doc.save(bio); bio.seek(0)
    return bio.getvalue()


def _h1(doc, text, accent):
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run(text)
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(accent)


def _merge_row(cells):
    a = cells[0]
    for c in cells[1:]:
        a = a.merge(c)


def _set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = w


# =========================================================================
# PDF
# =========================================================================
def _styles():
    ss = getSampleStyleSheet()
    out = {}
    out["Title"] = ParagraphStyle("CTitle", parent=ss["Normal"], fontName=F_BOLD, fontSize=20, leading=24,
                                  alignment=TA_CENTER, textColor=colors.HexColor("#2E3B4E"), spaceAfter=4)
    out["Sub"] = ParagraphStyle("CSub", parent=ss["Normal"], fontName=F_REG, fontSize=12, leading=15,
                                alignment=TA_CENTER, textColor=colors.HexColor("#444444"), spaceAfter=2)
    out["Who"] = ParagraphStyle("CWho", parent=ss["Normal"], fontName=F_REG, fontSize=9.5, leading=12,
                                alignment=TA_CENTER, textColor=colors.HexColor("#777777"), spaceAfter=10)
    out["Body"] = ParagraphStyle("CBody", parent=ss["Normal"], fontName=F_REG, fontSize=10, leading=14, spaceAfter=6)
    out["Cell"] = ParagraphStyle("CCell", parent=ss["Normal"], fontName=F_REG, fontSize=9.5, leading=12)
    out["CellB"] = ParagraphStyle("CCellB", parent=ss["Normal"], fontName=F_BOLD, fontSize=9.5, leading=12)
    out["CellW"] = ParagraphStyle("CCellW", parent=ss["Normal"], fontName=F_BOLD, fontSize=12, leading=15, textColor=colors.white)
    out["H1"] = ParagraphStyle("CH1", parent=ss["Normal"], fontName=F_BOLD, fontSize=13, leading=16,
                               textColor=colors.HexColor("#2E3B4E"), spaceBefore=10, spaceAfter=5)
    return out


def generate_conclusion_pdf_bytes(report: dict) -> bytes:
    lang = get_language(report)
    S = _styles()
    accent = report.get("theme_hex") or "#2E3B4E"
    for k in ("Title", "H1"):
        S[k].textColor = colors.HexColor(accent)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=16 * mm,
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            title=f"{report.get('client','')} Executive Conclusion")
    story = []
    story.append(Paragraph(EC.title(lang), S["Title"]))
    story.append(Paragraph(_safe(EC.subtitle(report, lang)), S["Sub"]))
    who = " · ".join([x for x in [_safe(report.get("client")).strip(), _safe(report.get("date")).strip()] if x])
    if who:
        story.append(Paragraph(who, S["Who"]))
    for para in EC.intro(report, lang):
        story.append(Paragraph(_safe(para), S["Body"]))
    story.append(Spacer(1, 6))

    # Metadata table
    meta = [[Paragraph(_safe(l), S["CellB"]), Paragraph(_safe(v), S["Cell"])] for l, v in EC.metadata_rows(report, lang)]
    mt = Table(meta, colWidths=[42 * mm, 128 * mm])
    mt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(mt); story.append(Spacer(1, 8))

    # Risk banner
    risk_label, risk_color, risk_desc = EC.risk(report, lang)
    rt = Table([[Paragraph(f'<font color="white"><b>{EC.section("risk_general", lang)}</b></font>', S["Cell"]), ""],
                [Paragraph(risk_label, S["CellW"]), Paragraph(_safe(risk_desc), S["Cell"])]],
               colWidths=[42 * mm, 128 * mm])
    rt.setStyle(TableStyle([
        ("SPAN", (0, 0), (1, 0)),
        ("BACKGROUND", (0, 0), (1, 0), colors.HexColor(accent)),
        ("BACKGROUND", (0, 1), (0, 1), colors.HexColor(risk_color)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(rt); story.append(Spacer(1, 6))

    pl = EC.posture_line(report, lang)
    if pl:
        pt = Table([[Paragraph(f"<b>{_safe(pl)}</b>", S["Cell"])]], colWidths=[170 * mm])
        pt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F3F5")),
                                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
                                ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story.append(pt)

    # Conclusion narrative
    story.append(Paragraph(EC.section("conclusion", lang), S["H1"]))
    for para in EC.narrative(report, lang):
        story.append(Paragraph(_safe(para), S["Body"]))

    # Impact table
    story.append(Paragraph(EC.section("impact", lang), S["H1"]))
    ih = EC.impact_header(lang)
    idata = [[Paragraph(ih[0], S["CellB"]), Paragraph(ih[1], S["CellB"])]]
    for area, impact in EC.impact_rows(lang):
        idata.append([Paragraph(_safe(area), S["CellB"]), Paragraph(_safe(impact), S["Cell"])])
    it = Table(idata, colWidths=[46 * mm, 124 * mm])
    it.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(it)

    # Remediation directions
    story.append(Paragraph(EC.section("remediation", lang), S["H1"]))
    story.append(Paragraph(_safe(EC.remediation_intro(report, lang)), S["Body"]))
    dh = EC.directions_header(lang)
    ddata = [[Paragraph(h, S["CellB"]) for h in dh]]
    for idx, d in enumerate(EC.directions(report, lang), start=1):
        ddata.append([Paragraph(str(idx), S["Cell"]), Paragraph(_safe(d.get("direction")), S["Cell"]),
                      Paragraph(_safe(d.get("intervention")), S["Cell"]), Paragraph(_safe(d.get("urgency")), S["Cell"])])
    dt = Table(ddata, colWidths=[8 * mm, 76 * mm, 54 * mm, 32 * mm], repeatRows=1)
    dt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(dt)

    # Compliance — priority findings mapped to controls (table only)
    mrows = EC.compliance_mapping_rows(report, lang)
    if mrows:
        def _esc(x):
            return _safe(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(EC.section("mapping", lang), S["H1"]))
        story.append(Paragraph(_esc(EC.mapping_lead(lang)), S["Body"]))
        mh = EC.mapping_header(lang)
        mdata = [[Paragraph(_esc(h), S["CellB"]) for h in mh]]
        for row in mrows:
            mdata.append([Paragraph(_esc(row["title"]), S["Cell"]),
                          Paragraph(_esc(row["severity"]), S["Cell"]),
                          Paragraph(_esc(row["domain"]), S["Cell"])])
        mt = Table(mdata, colWidths=[86 * mm, 26 * mm, 58 * mm], repeatRows=1)
        mt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F7")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(mt)

    doc.build(story)
    return buf.getvalue()
